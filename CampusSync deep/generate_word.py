import os
import subprocess
import sys
import base64

# Automatically install required packages
def install_packages():
    try:
        import docx
        import requests
    except ImportError:
        print("Installing required packages (python-docx, requests)...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "requests"])

install_packages()

import requests
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def get_mermaid_url(code):
    encoded = base64.b64encode(code.strip().encode('utf-8')).decode('utf-8')
    # Using default theme for Word documents so it prints well on white paper
    return f'https://mermaid.ink/img/{encoded}?theme=default'

chart1 = '''graph TD
Splash[App Launch / Splash] --> AuthCheck{Logged In?}
AuthCheck -- No --> Login[Login Screen]
AuthCheck -- Yes --> RoleCheck{Check User Role}
Login --> GoogleAuth[Continue with Google]
GoogleAuth --> DBCheck{Exists in users collection?}
DBCheck -- Yes --> RoleCheck
DBCheck -- No --> RequestForm[Request Access Form + College Code]
RequestForm --> SubmitRequest[Send Request to Firestore]
SubmitRequest --> WaitingRoom[Waiting Room]
WaitingRoom -- Admin Approves --> RoleCheck
RoleCheck -- Principal --> PD[Principal Dashboard]
RoleCheck -- Admin --> AD[Admin Dashboard]
RoleCheck -- Student --> SD[Student Dashboard]
RoleCheck -- Teacher --> TD[Teacher Dashboard]'''

chart2 = '''graph TD
Dash[Admin / Principal Dashboard] --> F1(Verify Requests)
Dash --> F2(Manage Staff)
Dash --> F3(Fee Management)
Dash --> F4(Principal Overview)
F1 --> List[Fetch Pending Users for this College]
List --> Action1{Approve or Reject}
Action1 -- Approve --> StatusUpdate[Update Status to 'approved']
Action1 -- Reject --> Delete[Delete User Request]
F2 --> StaffList[View Teachers & Admins]
StaffList --> AddStaff[Assign Roles]
F3 --> Setup[Create Fee Types]
F3 --> Assign[Assign Fees to Classes]
F3 --> VerifyPayment[Approve Student UTR/Payments]
F4 --> Analytics[View College-wide Attendance Analytics]'''

chart3 = '''graph TD
Dash[Teacher Dashboard] --> F1(Timetable & Lectures)
Dash --> F2(Take Attendance)
Dash --> F3(Upload Resources)
Dash --> F4(Assignments)
F1 --> ViewClasses[View Assigned Subjects & Classes]
F2 --> SelectSubject[Select Subject]
SelectSubject --> StudentList[View List of Students in Class]
StudentList --> Mark[Toggle Present / Absent]
Mark --> Save[Save to Firestore]
F3 --> AddFile[Upload PDF / Document]
AddFile --> Tagging[Tag by Subject & Semester]
Tagging --> CloudStorage[Upload to Firebase Storage]
F4 --> CreateTask[Create Assignment & Due Date]'''

chart4 = '''graph TD
Dash[Student Dashboard] --> F1(View Attendance)
Dash --> F2(View Resources)
Dash --> F3(Pay Fees)
Dash --> F4(Assignments)
F1 --> Stats[See Personal Attendance %]
F2 --> Search[Search Notes by Subject]
Search --> Download[Download PDF]
F3 --> PendingFees[View Assigned Pending Fees]
PendingFees --> Pay[Upload UTR Number & Screenshot]
Pay --> WaitApproval[Wait for Admin to Verify Payment]
F4 --> ToDo[View Pending Assignments]'''

chart5 = '''graph LR
    Root[(Firestore Root)] --> UsersColl[users collection]
    Root --> CollegesColl[colleges collection]
    
    UsersColl --> UserDoc[Document: uid]
    UserDoc -.->|Fields: status, role, collegeId| UserData
    
    CollegesColl --> CollegeDoc[Document: collegeId]
    CollegeDoc --> StudentsColl[students]
    CollegeDoc --> TeachersColl[teachers]
    CollegeDoc --> AdminsColl[admins]
    CollegeDoc --> PrincipalsColl[principals]
    
    CollegeDoc --> ClassesColl[classes]
    ClassesColl --> ClassDoc[Document: classId]
    ClassDoc --> AttColl[attendance]
    
    CollegeDoc --> ResColl[resources]
    CollegeDoc --> AssignColl[assignments]
    CollegeDoc --> FeesColl[fees]'''

chart6 = '''sequenceDiagram
    participant App as Android App
    participant Auth as Firebase Auth
    participant Storage as Firebase Cloud Storage
    participant DB as Firestore NoSQL DB
    
    Note over App,DB: 1. User Creation Logic
    App->>Auth: Authenticate via Google
    Auth-->>App: Return User UID
    App->>DB: Batch Write (Root User + Role Subcollection)
    
    Note over App,DB: 2. File & Resource Saving Logic
    App->>Storage: Upload PDF / Image File
    Storage-->>App: Return Public Download URL
    App->>DB: Save Resource Metadata (URL, Tags, Uploader)
    
    Note over App,DB: 3. Attendance Saving Logic
    App->>DB: Set Document (Map: studentId -> Present/Absent)
    
    Note over App,DB: 4. Real-time Logic (Snapshot Listeners)
    loop Real-time Updates
        DB-->>App: Push Status Changes (e.g. Pending -> Approved)
    end'''

document = Document()

# Title
title = document.add_heading('CampusSync Project Report', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Helper function to add sections
def add_section(title_text, chart_code, report_text):
    heading = document.add_heading(title_text, level=1)
    
    # Download image
    url = get_mermaid_url(chart_code)
    try:
        print(f"Downloading chart for {title_text}...")
        response = requests.get(url)
        if response.status_code == 200:
            image_stream = BytesIO(response.content)
            # Add image, width 6 inches fits well on A4/Letter
            doc_img = document.add_picture(image_stream, width=Inches(6.0))
            # Center the image
            last_paragraph = document.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            document.add_paragraph("[Failed to load flowchart image]")
    except Exception as e:
        document.add_paragraph(f"[Error downloading image: {e}]")
        
    document.add_heading('Report & Explanation', level=3)
    p = document.add_paragraph(report_text)
    
    # Add page break after each section except the last one ideally, but fine to just space it
    document.add_page_break()

# Sections
add_section(
    '1. Master App Flow (Overall Architecture)', 
    chart1, 
    "This flow demonstrates the bird's-eye view of CampusSync. It highlights the multi-tenancy entry point. When a new user logs in via Google, the system checks if they exist in the database. If not, they are directed to the Request Access Form where they must provide a unique College Code. They are then placed in a waiting room until their respective college Admin approves them. Once approved, they are routed to their role-specific dashboard."
)

add_section(
    '2. Principal & Admin Flow (Management & Verification)', 
    chart2, 
    "The Admin/Principal workflow handles high-level management. The most critical aspect is Verify Requests, where the system queries only those pending users whose collegeId matches the Admin's. The Admin can approve (updating the user's status to 'approved' in Firestore) or reject (deleting the user from the database). Additionally, Admins can create fee structures, assign them to specific classes, and verify UTR numbers for payments made by students."
)

add_section(
    '3. Teacher Flow (Attendance & Academics)', 
    chart3, 
    "Teachers are responsible for daily academic operations. They can select their assigned subjects and classes to Take Attendance (instantly saved and synced to Firestore). They can also Upload Resources (like notes and PDFs), tagging them properly by subject and semester for students to easily locate them in the cloud. Finally, they can create assignments and track which students have completed them."
)

add_section(
    '4. Student Flow (Learning & Fees)', 
    chart4, 
    "The Student dashboard provides read-only access to academic data and read-write access to fee payments. Students can monitor their real-time attendance statistics updated by teachers. They can search the digital library to download PDF notes uploaded by their instructors. For administration, students can view their pending fees, upload UPI payment screenshots along with the UTR number, and wait for the Admin to verify and clear their dues."
)

add_section(
    '5. Database Architecture (Firestore NoSQL Map)', 
    chart5, 
    "The Firebase Firestore database follows a multi-tenant hierarchy. The root 'users' collection tracks global identities, their active 'status', and their associated 'collegeId'. All primary data is nested under the 'colleges' collection. Within a specific college document, data is categorized into role-specific collections (students, teachers, admins, principals) and operational collections (classes, resources, assignments, fees). This ensures that data from one college is completely isolated from another."
)

add_section(
    '6. Core Logic & Data Saving Mechanisms', 
    chart6, 
    "This section outlines the internal business logic and how data is processed and saved in the cloud:\n\n"
    "1. Authentication & Batch Writing: When a user signs in, the app receives a UID from Firebase Auth. To ensure data consistency, the app uses 'Firestore Batches' to simultaneously create the user in the root 'users' collection and in their specific role subcollection. If either fails, the entire operation is rolled back.\n\n"
    "2. File Storage Logic: When a teacher uploads a resource or a student uploads a payment screenshot, the physical file is uploaded to Firebase Cloud Storage. The app waits for a secure 'Download URL' from Storage, and then saves that URL alongside metadata (tags, uploader name, timestamps) into the Firestore database.\n\n"
    "3. Attendance Logic: Instead of creating a separate document for every student per class, attendance is efficiently saved as a HashMap (key-value pair) within a single document for that lecture (e.g., studentId -> 'present' / 'absent').\n\n"
    "4. Real-time Sync Logic: The app heavily utilizes 'Snapshot Listeners'. Instead of manually refreshing the app, any changes in the database (like an Admin approving a student's pending request) are instantly pushed to the Android client, triggering automatic UI updates and navigation."
)

filename = 'CampusSync_Project_Report.docx'
document.save(filename)
print(f"Word document saved successfully as {filename}")
